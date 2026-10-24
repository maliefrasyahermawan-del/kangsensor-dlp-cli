"""Watchdog-backed file sensor — scan on create/modify events."""

from __future__ import annotations

import logging
import os
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import TYPE_CHECKING, Callable

from watchdog.events import FileSystemEvent, FileSystemEventHandler
from watchdog.observers import Observer

from kangsensor.sensors.base import SensorBase

if TYPE_CHECKING:
    from kangsensor.core.engine import Finding, RuleEngine

logger = logging.getLogger(__name__)

WATCHED_EXTENSIONS = {
    ".txt", ".csv", ".json", ".xml", ".log", ".yaml", ".yml",
    ".ini", ".cfg", ".conf", ".env",
    ".xlsx", ".xls", ".docx", ".doc", ".pdf",
    ".sql", ".py", ".js", ".ts", ".html", ".htm",
}

POOL_SIZE = 4


def _extract_text(path: Path) -> str | None:
    """Read text from a file, handling Office and PDF formats."""
    suffix = path.suffix.lower()
    try:
        if suffix in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            parts: list[str] = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    parts.append(" ".join(str(c) for c in row if c is not None))
            return "\n".join(parts)

        if suffix in (".docx", ".doc"):
            import docx
            doc = docx.Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs)

        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            return "\n".join(
                page.extract_text() or "" for page in reader.pages
            )

        with open(path, "rb") as fh:
            raw = fh.read()
        # Skip binary
        if len(raw) > 0 and raw.count(b"\x00") / len(raw) > 0.30:
            return None
        return raw.decode("utf-8", errors="replace")

    except Exception as exc:
        logger.debug("Cannot extract text from %s: %s", path, exc)
        return None


class _Handler(FileSystemEventHandler):
    def __init__(self, engine: RuleEngine, callback: Callable[[list[Finding]], None]) -> None:
        super().__init__()
        self._engine = engine
        self._callback = callback
        self._pool = ThreadPoolExecutor(max_workers=POOL_SIZE, thread_name_prefix="ks-file")

    def on_created(self, event: FileSystemEvent) -> None:
        self._handle(event)

    def on_modified(self, event: FileSystemEvent) -> None:
        self._handle(event)

    def _handle(self, event: FileSystemEvent) -> None:
        if event.is_directory:
            return
        path = Path(str(event.src_path))
        if path.suffix.lower() not in WATCHED_EXTENSIONS:
            return
        self._pool.submit(self._scan, path)

    def _scan(self, path: Path) -> None:
        text = _extract_text(path)
        if text is None:
            return
        findings = self._engine.evaluate(text, file_path=str(path))
        if findings:
            self._callback(findings)

    def shutdown(self) -> None:
        self._pool.shutdown(wait=False)


class FileSensor(SensorBase):
    """Monitor one or more directories for sensitive data in new/modified files."""

    def __init__(
        self,
        watch_paths: list[str],
        engine: RuleEngine,
        on_finding: Callable[[list[Finding]], None],
    ) -> None:
        super().__init__("file")
        self._paths = [Path(p).expanduser() for p in watch_paths]
        self._engine = engine
        self._on_finding = on_finding
        self._observer: Observer | None = None

    def _run(self) -> None:
        handler = _Handler(self._engine, self._on_finding)
        self._observer = Observer()
        for p in self._paths:
            if p.exists():
                self._observer.schedule(handler, str(p), recursive=True)
                logger.info("FileSensor watching: %s", p)
            else:
                logger.warning("Watch path does not exist: %s", p)

        self._observer.start()
        self._stop_event.wait()
        self._observer.stop()
        self._observer.join()
        handler.shutdown()

    def scan_now(self, target: str | Path) -> list[Finding]:
        """Synchronous one-shot scan of a file or directory."""
        target = Path(target).expanduser()
        results: list[Finding] = []
        paths_to_scan: list[Path] = []

        if target.is_file():
            paths_to_scan = [target]
        elif target.is_dir():
            for root, _, files in os.walk(target):
                for fname in files:
                    fpath = Path(root) / fname
                    if fpath.suffix.lower() in WATCHED_EXTENSIONS:
                        paths_to_scan.append(fpath)
        else:
            logger.error("Target not found: %s", target)
            return results

        with ThreadPoolExecutor(max_workers=POOL_SIZE) as pool:
            futures = [pool.submit(self._scan_file, p) for p in paths_to_scan]
            for f in futures:
                results.extend(f.result())

        return results

    def _scan_file(self, path: Path) -> list[Finding]:
        text = _extract_text(path)
        if text is None:
            return []
        return self._engine.evaluate(text, file_path=str(path))
